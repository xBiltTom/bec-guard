"""Controlled PDF/XLSX representations of the same privacy-limited report."""

from datetime import UTC, datetime
from html import escape
from io import BytesIO
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.reporting import OperationalReport

DISCLAIMER = "Los scores representan una evaluación del sistema y no probabilidades de fraude."
LABELS = {
    "LOW": "Bajo",
    "MEDIUM": "Medio",
    "HIGH": "Alto",
    "CRITICAL": "Crítico",
    "OPEN": "Abierta",
    "UNDER_REVIEW": "En revisión",
    "CLOSED": "Cerrada",
    "CONFIRMED_BEC": "BEC confirmado",
    "FALSE_POSITIVE": "Falso positivo",
    "EXECUTIVE_IMPERSONATION": "Suplantación ejecutiva",
    "VENDOR_FRAUD": "Fraude de proveedor",
    "BANK_ACCOUNT_CHANGE": "Cambio de cuenta",
    "SENSITIVE_DATA_REQUEST": "Solicitud de datos",
    "OTHER": "Otro patrón BEC",
    "open": "Abiertas",
    "under_review": "En revisión",
    "closed": "Cerradas",
    "confirmed_bec": "BEC confirmados",
    "false_positive": "Falsos positivos",
}


def label(value) -> str:
    return LABELS.get(value, str(value)) if value is not None else "Sin clasificar"


def summary_rows(report: OperationalReport) -> list[list]:
    s = report.summary
    return [
        ["Correos recibidos", s["emails_total"]],
        ["Análisis completados", s["analyses_completed"]],
        ["Errores de análisis", s["analyses_error"]],
        *[[label(k), v] for k, v in s["alerts"].items()],
        *[[label(k), v] for k, v in s["verdicts"].items()],
    ]


def xlsx_bytes(report: OperationalReport) -> bytes:
    book = Workbook()
    book.remove(book.active)
    tables = [
        ("Resumen", ["Indicador", "Cantidad"], summary_rows(report), [42, 20]),
        (
            "Distribución de riesgo",
            ["Riesgo", "Cantidad"],
            [[label(k), v] for k, v in report.summary["risk_distribution"].items()],
            [32, 20],
        ),
        (
            "Tipos BEC",
            ["Tipo identificado", "Cantidad"],
            [[label(t["bec_type"]), t["count"]] for t in report.summary["bec_type_distribution"]],
            [38, 20],
        ),
        (
            "Alertas",
            [
                "Alerta ID",
                "Creada (UTC)",
                "Score",
                "Riesgo",
                "Tipo BEC",
                "Estado",
                "Veredicto",
                "Responsable",
            ],
            [
                [
                    str(a.id),
                    a.created_at.astimezone(UTC).replace(tzinfo=None),
                    a.risk_score,
                    label(a.risk_level),
                    label(a.bec_type),
                    label(a.status),
                    label(a.verdict) if a.verdict else "Sin decisión",
                    a.assigned_to_name or "Sin asignar",
                ]
                for a in report.alerts
            ],
            [39, 23, 10, 14, 29, 19, 21, 30],
        ),
    ]
    for title, headers, rows, widths in tables:
        sheet = book.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.append(["BEC Guard | Reporte operativo"])
        sheet.append([report.period])
        sheet.append(["Generado (UTC)", report.generated_at.astimezone(UTC).replace(tzinfo=None)])
        sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{max(5, sheet.max_row)}"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers)))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(headers)))
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[5].height = 28
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Never Excel formulas from names.
                cell.font = Font(name="Calibri", size=11, color="233747")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                if cell.row == 5:
                    cell.fill = PatternFill("solid", fgColor="172E3F")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                elif cell.row > 5:
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F1F5F7")
                    sheet.row_dimensions[cell.row].height = 34 if title == "Alertas" else 25
        for row in sheet.iter_rows(min_row=6):
            lines = 1
            for cell in row:
                value = cell.value
                text = (
                    value.strftime("%Y-%m-%d %H:%M:%S")
                    if isinstance(value, datetime)
                    else str(value or "")
                )
                width = widths[cell.column - 1]
                lines = max(
                    lines,
                    sum(max(1, ceil(len(part) / max(1, width - 2))) for part in text.split("\n")),
                )
            sheet.row_dimensions[row[0].row].height = max(
                34 if title == "Alertas" else 25, lines * 15 + 6
            )
        sheet["A1"].font = Font(name="Calibri", size=16, bold=True, color="172E3F")
        sheet.page_setup.orientation = "landscape" if title == "Alertas" else "portrait"
        sheet.print_title_rows = "1:5"
    book["Resumen"].append([])
    book["Resumen"].append([DISCLAIMER])
    final_row = book["Resumen"].max_row
    book["Resumen"].merge_cells(
        start_row=final_row, start_column=1, end_row=final_row, end_column=2
    )
    book["Resumen"].row_dimensions[final_row].height = 38
    book["Resumen"].cell(final_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    output = BytesIO()
    try:
        book.save(output)
        return output.getvalue()
    finally:
        book.close()
        output.close()


def pdf_html(report: OperationalReport) -> str:
    def table(headers, rows, css=""):
        head = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
        body = "".join(
            "<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>" for row in rows
        )
        if not body:
            body = f'<tr><td colspan="{len(headers)}">Sin registros en el periodo.</td></tr>'
        return f'<table class="{css}"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'

    s = report.summary
    sections = "<h2>Resumen ejecutivo</h2>" + table(
        ["Indicador", "Cantidad"], summary_rows(report)[:3]
    )
    sections += '<div class="paired"><section><h2>Distribución de riesgo</h2>' + table(
        ["Riesgo", "Cantidad"], [[label(k), v] for k, v in s["risk_distribution"].items()]
    )
    sections += "</section><section><h2>Tipos BEC</h2>" + table(
        ["Tipo identificado", "Cantidad"],
        [[label(t["bec_type"]), t["count"]] for t in s["bec_type_distribution"]],
    )
    sections += '</section></div><div class="paired"><section><h2>Estado de alertas</h2>' + table(
        ["Estado actual", "Cantidad"], [[label(k), v] for k, v in s["alerts"].items()]
    )
    sections += "</section><section><h2>Veredictos</h2>" + table(
        ["Decisión humana", "Cantidad"], [[label(k), v] for k, v in s["verdicts"].items()]
    )
    sections += (
        '</section></div><section class="alert-list"><h2>Alertas del periodo</h2>'
        + table(
            [
                "Alerta / creada UTC",
                "Score / riesgo",
                "Tipo BEC",
                "Estado / veredicto",
                "Responsable",
            ],
            [
                [
                    f"{a.id}\n{a.created_at.astimezone(UTC):%Y-%m-%d %H:%M}",
                    f"{a.risk_score} / {label(a.risk_level)}",
                    label(a.bec_type),
                    f"{label(a.status)} / {label(a.verdict) if a.verdict else 'Sin decisión'}",
                    a.assigned_to_name or "Sin asignar",
                ]
                for a in report.alerts
            ],
            "alerts",
        )
        + "</section>"
    )
    return (
        """<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>BEC Guard - Reporte operativo</title>
<style>

@page {
  size: A4; margin: 16mm 15mm 21mm;
  @bottom-left { content: "BEC Guard | Reporte operativo"; font-size: 8pt; color: #586c7a; }
  @bottom-right { content: counter(page) " / " counter(pages); font-size: 8pt; color: #586c7a; }
}
body { font-family: "DejaVu Sans", sans-serif; font-size: 9pt; color: #233747; line-height: 1.45; }
header { border-bottom: 3px solid #126b70; padding-bottom: 12pt; margin-bottom: 18pt; }
header strong { font-size: 18pt; color: #172e3f; }
h1 { font-size: 15pt; font-weight: normal; margin: 6pt 0; }
header p { font-size: 8pt; margin: 3pt 0; color: #586c7a; }
h2 { font-size: 11pt; color: #172e3f; margin: 17pt 0 7pt; break-after: avoid; }
table { width: 100%; border-collapse: collapse; table-layout: fixed; }
th, td {
  text-align: left; vertical-align: top; padding: 6pt 8pt;
  border-bottom: 1px solid #d5dfe5; overflow-wrap: anywhere; white-space: pre-wrap;
}
th { background: #e7eff2; font-size: 8pt; }
th:last-child, td:last-child { width: 29%; }
tr { break-inside: avoid; }
thead { display: table-header-group; }
tbody tr:nth-child(even) { background: #f6f8fa; }
.paired { display: table; width: 100%; table-layout: fixed; break-inside: avoid; }
.paired > section { display: table-cell; width: 50%; vertical-align: top; }
.paired > section:first-child { padding-right: 10pt; }
.paired > section:last-child { padding-left: 10pt; }
.paired th:last-child, .paired td:last-child { width: 25%; }
.alert-list { break-before: page; }
.alerts { font-size: 7pt; }
.alerts th, .alerts td { padding: 5pt 4pt; }
.alerts th:first-child { width: 29%; }
.alerts th:last-child, .alerts td:last-child { width: 16%; }
.disclaimer { margin-top: 18pt; font-size: 8pt; color: #586c7a; }
</style></head>
<body><header><strong>BEC Guard</strong>
<h1>Reporte operativo</h1>
<p>Periodo: """
        + escape(report.period)
        + "</p><p>Generado: "
        + report.generated_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")
        + "</p></header>"
        + sections
        + '<p class="disclaimer">'
        + escape(DISCLAIMER)
        + "</p></body></html>"
    )


def pdf_bytes(report: OperationalReport) -> bytes:
    from weasyprint import HTML

    def deny_fetch(*_args, **_kwargs):
        raise ValueError("Report templates do not fetch external resources")

    return HTML(string=pdf_html(report), url_fetcher=deny_fetch).write_pdf()
