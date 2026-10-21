from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Role, Vendor
from app.services.alerts import ensure_alert_for_analysis
from app.services.report_formats import pdf_bytes, pdf_html, xlsx_bytes
from app.services.reporting import gather_report

pytestmark = pytest.mark.integration


@pytest.fixture
def report_setup(db_session, account, alert_setup):
    email, assessment, headers = alert_setup
    analysis = assessment(85)
    alert = ensure_alert_for_analysis(db_session, analysis, actor=account)
    account.role = Role.ADMIN
    account.name = '=HYPERLINK("https://untrusted.example","Name")'
    db_session.add(
        Vendor(
            name="Report vendor",
            domain="report-vendor.example",
            bank_account_reference="PRIVATE-BANK-REPORT",
        )
    )
    email.ingested_at = datetime(2026, 10, 20, 23, 59, 59, tzinfo=UTC)
    analysis.analyzed_at = datetime(2026, 10, 21, 0, 0, tzinfo=UTC)
    alert.created_at = datetime(2026, 10, 21, 23, 59, 59, tzinfo=UTC)
    alert.status, alert.assigned_to = "UNDER_REVIEW", account.id
    db_session.commit()
    return headers, alert, email


@pytest.mark.parametrize("role", list(Role))
@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_report_rbac_and_http(api_client, db_session, account, report_setup, role, extension):
    headers, _, _ = report_setup
    account.role = role
    db_session.commit()
    result = api_client.get(f"/api/reports/operational.{extension}", headers=headers)
    if role == Role.ANALYST:
        assert result.status_code == 403
        return
    assert result.status_code == 200
    assert result.headers["content-disposition"].startswith(
        'attachment; filename="bec-guard-operational-'
    )
    assert result.headers["content-disposition"].endswith(f'.{extension}"')
    assert result.headers["cache-control"] == "no-store"
    assert result.content.startswith(b"%PDF-" if extension == "pdf" else b"PK")
    assert result.headers["content-type"] == (
        "application/pdf"
        if extension == "pdf"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert api_client.get(f"/api/reports/operational.{extension}").status_code == 401
    assert (
        api_client.get(
            f"/api/reports/operational.{extension}", headers={"X-BEC-Automation-Key": "x" * 48}
        ).status_code
        == 401
    )


@pytest.mark.parametrize(
    "query", ["from_date=2026-10-22&to_date=2026-10-21", "from_date=invalid", "to_date=2026-02-30"]
)
@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_report_invalid_dates(api_client, report_setup, query, extension):
    headers, _, _ = report_setup
    assert (
        api_client.get(f"/api/reports/operational.{extension}?{query}", headers=headers).status_code
        == 400
    )


def test_report_dates_and_aggregation(db_session, report_setup):
    _, alert, _ = report_setup
    all_data = gather_report(db_session)
    assert all_data.summary["emails_total"] == all_data.summary["analyses_completed"] == 1
    assert all_data.summary["analyses_error"] == 0
    assert all_data.summary["risk_distribution"]["CRITICAL"] == 1
    assert all_data.summary["alerts"]["under_review"] == 1
    assert all_data.alerts[0].id == alert.id
    assert all_data.alerts[0].assigned_to_name.startswith("=HYPERLINK")
    before = gather_report(db_session, date(2026, 10, 20), date(2026, 10, 20))
    assert before.summary["emails_total"] == 1 and before.summary["analyses_completed"] == 0
    assert not before.alerts
    day = gather_report(db_session, date(2026, 10, 21), date(2026, 10, 21))
    assert day.summary["emails_total"] == 0 and day.summary["analyses_completed"] == 1
    assert len(day.alerts) == 1
    after = gather_report(db_session, date(2026, 10, 22), None)
    assert after.summary["analyses_completed"] == 0 and not after.alerts
    assert gather_report(db_session, None, date.max).summary["emails_total"] == 1
    assert gather_report(db_session, None, date(2026, 10, 20)).summary["emails_total"] == 1
    with pytest.raises(ValueError):
        gather_report(db_session, date(2026, 10, 22), date(2026, 10, 21))


def test_xlsx_structure_types_privacy_and_formula_safety(db_session, account, report_setup):
    _, _, email = report_setup
    report = gather_report(db_session)
    workbook = load_workbook(BytesIO(xlsx_bytes(report)))
    assert workbook.sheetnames == ["Resumen", "Distribución de riesgo", "Tipos BEC", "Alertas"]
    assert workbook["Resumen"]["B6"].value == 1
    assert workbook["Alertas"]["C6"].value == 85
    assert isinstance(workbook["Alertas"]["B6"].value, datetime)
    assert workbook["Alertas"]["H6"].value.startswith("=HYPERLINK")
    assert workbook["Alertas"]["H6"].data_type == "s"
    assert workbook["Alertas"].freeze_panes == "A6"
    assert workbook["Alertas"].auto_filter.ref == "A5:H6"
    values = repr([[cell.value for row in sheet for cell in row] for sheet in workbook])
    for forbidden in (
        email.body_text,
        email.subject,
        email.sender_email,
        email.reply_to,
        email.authentication_results,
        "PRIVATE-BANK-REPORT",
        account.password_hash,
        "ground_truth",
        "expected_bec_type",
        "access_token",
    ):
        assert forbidden not in values
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_pdf_controlled_html_and_no_external_fetch(db_session, report_setup, account, monkeypatch):
    _, _, email = report_setup
    account.name = '<img src="https://untrusted.example"> & unsafe'
    db_session.commit()
    report = gather_report(db_session)
    html = pdf_html(report)
    assert "<img src=" not in html and "&lt;img src=" in html
    assert "Reporte operativo" in html and "probabilidades de fraude" in html
    for forbidden in (
        email.body_text,
        email.subject,
        email.sender_email,
        "PRIVATE-BANK-REPORT",
        "ground_truth",
        "expected_bec_type",
    ):
        assert forbidden not in html
    # The renderer needs no URL fetches: fonts come from the installed font system.
    import weasyprint

    real_html = weasyprint.HTML

    def controlled_html(*args, **kwargs):
        with pytest.raises(ValueError, match="external resources"):
            kwargs["url_fetcher"]("https://untrusted.example/image.png")
        return real_html(*args, **kwargs)

    monkeypatch.setattr(weasyprint, "HTML", controlled_html)
    assert pdf_bytes(report).startswith(b"%PDF-")


def test_empty_report(db_session):
    report = gather_report(db_session)
    assert report.summary["emails_total"] == report.summary["analyses_completed"] == 0
    assert not report.alerts and not report.summary["bec_type_distribution"]
    assert pdf_bytes(report).startswith(b"%PDF-")
    workbook = load_workbook(BytesIO(xlsx_bytes(report)))
    assert workbook["Resumen"]["B6"].value == 0
    assert workbook["Alertas"].max_row == 5
    workbook.close()


def test_error_count_is_not_low(db_session, account, alert_setup):
    _, assessment, _ = alert_setup
    assessment(100, status="ERROR")
    report = gather_report(db_session)
    assert report.summary["analyses_error"] == 1 and report.summary["analyses_completed"] == 0
    assert sum(report.summary["risk_distribution"].values()) == 0
    assert not report.alerts
